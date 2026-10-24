"""Lessons from running the MCP on a real multi-floor project: gates, protected bases, intake,
sanitising heavy exports, table-driven placement and verification, lengths, quantities, recovery."""

import re

import ezdxf
import openpyxl
import pytest

from acad_electrical_mcp import plan as plans
from acad_electrical_mcp import server, tables
from acad_electrical_mcp.bases import BaseRegistry
from acad_electrical_mcp.live import LiveSession, PlainCom
from acad_electrical_mcp.model import Store
from acad_electrical_mcp.oplog import OpLog
from fake_acad import App, Block, Ent, Layer


@pytest.fixture()
def s(tmp_path, monkeypatch):
    ws = tmp_path / "ws"
    monkeypatch.setattr(server.CFG, "workspace", ws)
    monkeypatch.setattr(server.CFG, "require_plan", False)
    monkeypatch.setattr(server, "STORE", Store(ws))
    monkeypatch.setattr(server, "PLANS", plans.PlanStore(ws))
    monkeypatch.setattr(server, "BASES", BaseRegistry(ws))
    monkeypatch.setattr(server, "OPS", OpLog(ws, lambda: server.PLANS.active()))
    return server


@pytest.fixture()
def live(s, tmp_path, monkeypatch):
    app = App()
    sess = LiveSession(PlainCom(app), tmp_path / "st")
    monkeypatch.setattr(server, "_LIVE", sess)
    sess.is_protected = s.BASES.is_protected
    sess.oplog = s.OPS.append
    real = tmp_path / "Drawing1.dwg"
    real.write_bytes(b"AC1032")
    app.ActiveDocument.FullName = str(real)
    return sess, app


def plan_with_floors(s, floors=("GF", "1F")):
    s.plan_define("proj", "Floors", "Draw the floors one at a time.", ["each floor confirmed"],
                  floors=list(floors))
    return s.plan_approve("proj", user_confirmed=True)


# ------------------------------------------------------------------ floor gates
def test_one_floor_at_a_time_and_confirmation(s):
    plan_with_floors(s)
    with pytest.raises(ValueError, match="user_confirmed"):
        s.plan_approve_floor("proj", "GF", user_confirmed=False)
    s.plan_approve_floor("proj", "GF", user_confirmed=True, note="start with GF")
    with pytest.raises(ValueError, match="still open"):
        s.plan_approve_floor("proj", "1F", user_confirmed=True)
    with pytest.raises(ValueError, match="never approved"):
        s.plan_confirm_floor("proj", "1F", user_confirmed=True)
    out = s.plan_confirm_floor("proj", "GF", user_confirmed=True, note="checked")
    assert out["floors"]["GF"]["confirmed"] is True
    out = s.plan_approve_floor("proj", "1F", user_confirmed=True)
    assert "1F OPEN" in out["checklist"] and out["still_open"]["floors_awaiting_confirmation"] == ["1F"]


def test_gate_messages_for_floors():
    plan = plans.define("p", "t", "g", ["c"], floors=["GF", "1F"])
    assert "no approved plan" in plans.gate_check(plan)
    plan["goal"]["approved"] = True
    assert "no floor is approved" in plans.gate_check(plan)
    plans.approve_floor(plan, "GF")
    assert plans.gate_check(plan) is None and plans.gate_check(plan, "GF") is None
    assert "not approved for drawing" in plans.gate_check(plan, "1F")
    plans.confirm_floor(plan, "GF")
    assert "already confirmed" in plans.gate_check(plan, "GF")
    plans.approve_floor(plan, "1F")
    plans.approve_floor(plan, "GF", force=True)
    assert "several floors are open" in plans.gate_check(plan)


def test_strict_mode_follows_the_open_floor(s, live, monkeypatch):
    sess, app = live
    monkeypatch.setattr(server.CFG, "require_plan", True)
    plan_with_floors(s)
    with pytest.raises(ValueError, match="GF is not approved"):
        s.live_place("GF", "socket", 10, 10)
    s.plan_approve_floor("proj", "GF", user_confirmed=True)
    assert s.live_place("GF", "socket", 10, 10)["created"]
    with pytest.raises(ValueError, match="1F is not approved"):
        s.live_place("1F", "socket", 20, 20)
    assert s.live_move("GF-SKT-01", dx=5)["x"] == 15.0  # active floor used when no floor is given
    s.plan_confirm_floor("proj", "GF", user_confirmed=True)
    with pytest.raises(ValueError, match="no floor is approved"):
        s.live_move("GF-SKT-01", dx=5)  # confirmed floors are locked
    monkeypatch.setattr(server.CFG, "require_plan", False)


def test_open_items_and_impact_map(s):
    plan_with_floors(s)
    s.plan_note("proj", "Floor box overlaps the A/C cassette", kind="open")
    sh = s.plan_show("proj")
    assert sh["still_open"]["open_items"] == [{"number": 0, "text": "Floor box overlaps the A/C cassette"}]
    with pytest.raises(ValueError, match="No open item"):
        s.plan_resolve("proj", 5)
    assert s.plan_resolve("proj", 0, how="moved 700 mm")["still_open"]["open_items"] == []
    imp = s.plan_impact("demand factors")
    assert {"cable sizes", "single line diagrams", "BOQ", "report"} <= set(imp["affected"])
    assert imp["cad_drawings_affected"] is False
    assert s.plan_impact("circuits")["cad_drawings_affected"] is True
    with pytest.raises(ValueError, match="Unknown change"):
        s.plan_impact("banana")


def test_new_template_and_prompts(s):
    names = {t["template"]: t for t in s.plan_templates()["templates"]}
    assert "floor_set_workflow" in names and names["floor_set_workflow"]["not_built"]
    out = s.plan_define("p2", "Set", "g", ["c"], template="floor_set_workflow", floors=["GF"])
    assert out["progress"]["total"] == 12
    import asyncio
    prompts = {p.name for p in asyncio.run(server.mcp.list_prompts())}
    assert {"start_project", "new_floor", "finish_floor_set", "plain_language"} <= prompts
    assert "ONE FLOOR AT A TIME" in server.INSTRUCTIONS and "BASE DRAWINGS ARE NEVER WRITTEN" in server.INSTRUCTIONS


# ------------------------------------------------------------------ protected bases
def make_dxf(path, layers=(("Wall", 5),)):
    doc = ezdxf.new("R2018")
    for lay, n in layers:
        if lay not in doc.layers:
            doc.layers.add(lay)
        for i in range(n):
            doc.modelspace().add_line((i, 0), (i, 1), dxfattribs={"layer": lay})
    doc.saveas(str(path))
    return str(path)


def test_base_copy_never_touches_the_base(s, tmp_path):
    base = make_dxf(tmp_path / "base.dxf")
    out = s.floor_from_base(base, "GF", str(tmp_path / "work"))
    assert out["working_copy"].endswith("GF_base.dxf") and out["identical_to_base"]
    assert out["base_unchanged"] is True and s.base_check()["ok"]
    with pytest.raises(ValueError, match="already exists"):
        s.floor_from_base(base, "GF", str(tmp_path / "work"))
    with pytest.raises(ValueError, match="registered base"):
        s.floor_from_base(base, "GF", str(tmp_path), name="base.dxf")
    with open(base, "ab") as fh:
        fh.write(b"\n0\n")
    chk = s.base_check()
    assert not chk["ok"] and len(chk["changed"]) == 1


def test_writes_and_saves_are_refused_on_a_base_drawing(s, live):
    sess, app = live
    s.base_register(app.ActiveDocument.FullName, role="architectural base")
    assert s.live_connect()["active_drawing_is_protected_base"] is True
    for call in (lambda: s.live_place("GF", "socket", 1, 1), lambda: s.live_add_text("TEXT", "x", 0, 0, 100),
                 lambda: s.live_save(), lambda: s.live_save(app.ActiveDocument.FullName)):
        with pytest.raises(ValueError, match="registered BASE"):
            call()
    assert s.live_inventory is not None and s.live_list_devices()["devices"] == []  # reading is fine
    app.ActiveDocument.FullName = "C:/x/GF_working.dwg"  # the user opened / saved a copy
    assert s.live_place("GF", "socket", 1, 1)["created"]
    assert s.live_save("C:/x/GF_final.dwg")["saved"].endswith("GF_final.dwg")


# ------------------------------------------------------------------ intake
def test_intake_scan_finds_floors_scale_and_heavy_files(s, tmp_path):
    folder = tmp_path / "inputs"
    (folder / "All Dialux Exports").mkdir(parents=True)
    for f, span in (("GF", 100), ("1F", 100), ("2F", 57), ("3F", 100), ("4F", 100), ("6F", 100)):
        doc = ezdxf.new("R2018")
        doc.units = 6
        doc.modelspace().add_line((0, 0), (span, span / 2))
        doc.saveas(str(folder / "All Dialux Exports" / f"floor {f} dialux export.dxf"))
    heavy = ezdxf.new("R2018")
    heavy.layers.add("DLX_OBJ")
    heavy.modelspace().add_line((0, 0), (1, 1), dxfattribs={"layer": "DLX_OBJ"})
    heavy.saveas(str(folder / "raw dialux export 5F.dxf"))
    (folder / "4F_Architectural_Base.dwg").write_bytes(b"AC1032" + b"\0" * 200)
    (folder / "Savindu.evo").write_bytes(b"PK" + b"\0" * 50)
    (folder / "ProjectToc.rsl").write_bytes(b"serialization::archive")
    wb = openpyxl.Workbook()
    wb.active.title = "Demand"
    wb.save(folder / "1F_Maximum_Demand.xlsx")
    (folder / "copy of Savindu.evo").write_bytes(b"PK" + b"\0" * 50)

    res = s.intake_scan(str(folder))
    kinds = res["by_kind"]
    assert kinds["dxf"] == 7 and kinds["dwg"] == 1 and kinds["excel"] == 1
    assert kinds["dialux_project"] == 2 and kinds["dialux_internal"] == 1
    by_file = {f["file"].split("/")[-1].split("\\")[-1]: f for f in res["files"]}
    assert by_file["4F_Architectural_Base.dwg"]["dwg_version"] == "2018 or newer"
    assert by_file["4F_Architectural_Base.dwg"]["floor"] == "4F" and by_file["1F_Maximum_Demand.xlsx"]["sheets"] == ["Demand"]
    assert "5F" in res["floors_found"]["dialux_export"] and not res["missing_floors"]
    assert s.intake_scan(str(folder), expected_floors=["GF", "1F", "7F"])["missing_floors"]["dialux_export"] == ["7F"]
    assert any("floor 2F dialux export.dxf" in f and "wrong scale" in f for f in res["scale_flags"])
    assert res["duplicates"] and any("Savindu" in x for x in res["duplicates"][0])
    text = " ".join(res["warnings"])
    assert "Heavy DIALux" in text and "DWG contents cannot be read" in text and ".evo is a DIALux project" in text
    assert ".rsl files" in text and res["ask_the_user"]
    with pytest.raises(ValueError, match="not found"):
        s.intake_scan(str(tmp_path / "nope"))


def test_floor_detection_from_names():
    from acad_electrical_mcp.intake import floor_of
    cases = {"4F_Architectural_Base": "4F", "floor gf new dialux export": "GF", "Ground Floor dialux export": "GF",
             "floor 1 dialux export": "1F", "6F_Electrical_Complete": "6F", "second floor plan": "2F",
             "Savindu": None, "Project": None}
    for name, want in cases.items():
        assert floor_of(name) == want, name


# ------------------------------------------------------------------ sanitising heavy DXF
def test_dxf_sanitize_makes_a_slim_scaled_copy(s, tmp_path):
    src = tmp_path / "raw.dxf"
    doc = ezdxf.new("R2018")
    doc.units = 6
    for lay in ("DLX_OBJ", "DLX_CALC", "Wall"):
        doc.layers.add(lay)
    m = doc.modelspace()
    for i in range(50):
        m.add_line((i, 0), (i, 1), dxfattribs={"layer": "DLX_OBJ"})
    m.add_line((0, 0), (10, 0), dxfattribs={"layer": "DLX_CALC"})
    m.add_line((0, 0), (10, 5), dxfattribs={"layer": "Wall"})
    doc.saveas(str(src))
    out = s.dxf_sanitize(str(src), str(tmp_path / "slim.dxf"), ["DLX_OBJ", "DLX_CALC"], from_units="m")
    assert out["erased_entities"] == 51 and out["source_unchanged"] and out["scale_applied"] == 1000
    assert out["before"]["entities"] == 52 and out["after"]["entities"] == 1
    assert out["after"]["extents"][2] == pytest.approx(10000) and "millimetres" in out["units_now"]
    s.base_register(str(src))
    for bad in ({"destination": str(src)}, {"destination": str(tmp_path / "slim.dxf")}):
        with pytest.raises(ValueError, match="registered base|already exists"):
            s.dxf_sanitize(str(src), strip_layers=["DLX_OBJ"], **bad)
    with pytest.raises(ValueError, match="scale or from_units"):
        s.dxf_sanitize(str(src), str(tmp_path / "x.dxf"), [], scale=2, from_units="m")
    (tmp_path / "a.dwg").write_bytes(b"AC1032")
    with pytest.raises(ValueError, match="DXF only"):
        s.dxf_sanitize(str(tmp_path / "a.dwg"), str(tmp_path / "x.dxf"), [])


# ------------------------------------------------------------------ placement table
ROOMS = [{"id": "GF-1", "name": "LAB", "bounds": [0, 0, 6000, 5000]},
         {"id": "GF-2", "name": "CLASS", "bounds": [6000, 0, 12000, 5000]}]


def filled_table(s, rows):
    path = s.workbook_create("proj", "placement_table", overwrite=True)["path"]
    wb = openpyxl.load_workbook(path)
    for r, vals in enumerate(rows, start=2):
        for c, v in enumerate(vals, start=1):
            wb["Data"].cell(row=r, column=c, value=v)
    wb.save(path)
    return path


def model_with_rooms(s, tmp_path):
    ref = tmp_path / "base.dxf"
    make_dxf(ref)
    s.prepare_floor_model("proj", "GF", str(ref), ROOMS)
    s.apply_changes("proj", "GF", s.plan_devices("proj", "GF", [{"type": "db", "x": 100, "y": 4900}])["changeset_id"])


TABLE = [
    ["GF", "LAB", "socket", 500, 400, 0, "GF-P01", "DB-GF", "GF-R1", ""],
    ["GF", "LAB", "socket", 2500, 400, 0, "GF-P01", "DB-GF", "GF-R2", ""],
    ["GF", "LAB", "data", 3000, 4500, 0, "GF-P01", "DB-GF", "GF-D1", ""],
    ["GF", "CLASS", "socket", 6500, 400, 0, "GF-P02", "DB-GF", "GF-R3", ""],
    ["GF", "CLASS", "ac", 9000, 4500, 0, "GF-AC1", "DB-GF", "GF-AC1", "ceiling cassette"],
    ["1F", "LAB", "socket", 500, 400, 0, "1F-P01", "DB-1F", "", ""],
]


def test_place_from_table_verify_and_batches(s, tmp_path):
    model_with_rooms(s, tmp_path)
    table = filled_table(s, TABLE)
    cs = s.place_from_table("proj", "GF", table, limit=3)
    assert cs["table_rows_for_floor"] == 5 and cs["rows_in_this_batch"] == 3
    assert cs["remaining_after_this_batch"] == 2 and cs["operations"] >= 3
    s.apply_changes("proj", "GF", cs["changeset_id"])
    cs2 = s.place_from_table("proj", "GF", table, offset=3, limit=3)
    assert cs2["remaining_after_this_batch"] == 0
    s.apply_changes("proj", "GF", cs2["changeset_id"])
    assert s.place_from_table("proj", "GF", table)["operations"] == 0  # idempotent
    v = s.verify_placement("proj", "GF", table)
    assert v["match"] and v["expected_total"] == 5 and v["actual_total"] == 5
    devs = {d["id"]: d for d in s.list_devices("proj", "GF")["devices"]}
    assert next(d for d in devs.values() if d["type"] == "ac")["room"] == "GF-2"
    # remove one device behind the table's back: the match table names it
    victim = next(i for i, d in devs.items() if d["type"] == "data")
    s.apply_changes("proj", "GF", s.plan_devices("proj", "GF", [], remove_ids=[victim])["changeset_id"])
    bad = s.verify_placement("proj", "GF", table)
    assert not bad["match"] and bad["mismatches"] == [
        {"room": "GF-1", "type": "data", "expected": 1, "actual": 0, "match": False}]


def test_placement_table_rejects_bad_input_and_copied_rooms(s, tmp_path):
    model_with_rooms(s, tmp_path)
    bad = filled_table(s, [["GF", "LAB", "toaster", "x", 1, 0, "", "", "", ""]])
    with pytest.raises(ValueError, match="error"):
        s.place_from_table("proj", "GF", bad)
    same = [["GF", room, t, x, y, 0, "", "", "", ""]
            for room, ox in (("LAB", 0), ("CLASS", 6000))
            for t, x, y in (("socket", 500, 400), ("socket", 2500, 400), ("data", 3000, 4500))]
    for r in same:
        r[3] += 6000 if r[1] == "CLASS" else 0
    path = filled_table(s, same)
    with pytest.raises(ValueError, match="identical device lists"):
        s.place_from_table("proj", "GF", path)
    assert s.place_from_table("proj", "GF", path, allow_identical=True)["operations"] > 0
    assert tables.identical_rooms([{**r, "room": "A", "type": "socket", "notes": "why"} for r in [{}, {}, {}]]) == []
    with pytest.raises(ValueError, match="No rows for floor"):
        s.place_from_table("proj", "GF", filled_table_other_floor(s))


def filled_table_other_floor(s):
    p = s.workbook_create("proj2", "placement_table")["path"]
    wb = openpyxl.load_workbook(p)
    for c, v in enumerate(["3F", "LAB", "socket", 1, 1, 0, "", "", "", ""], start=1):
        wb["Data"].cell(row=2, column=c, value=v)
    wb.save(p)
    return p


def test_exports_lengths_and_quantities_as_named_ranges(s, tmp_path):
    from conftest import build_floor
    ref = tmp_path / "ref.dxf"
    from acad_electrical_mcp.sample import ROOMS as SAMPLE_ROOMS
    from acad_electrical_mcp.sample import make_sample_floor
    make_sample_floor(str(ref))
    s.prepare_floor_model("demo", "4F", str(ref), SAMPLE_ROOMS)
    build_floor(s)
    s.plan_devices("demo", "4F", [{"type": "socket", "x": 5000, "y": 100, "circuit": "P9"}])  # circuit P9 has no route
    out = s.export_circuit_lengths("demo", "4F", factor=1.1)
    assert out["unit_to_m"] == 0.001 and out["named_range"] == "Circuit_Lengths_4F"
    wb = openpyxl.load_workbook(out["file"])
    assert "Circuit_Lengths_4F" in wb.defined_names
    rows = {r["circuit"]: r for r in out["rows"]}
    assert rows["L1"]["route_length_m"] > 0
    assert rows["L1"]["length_m_with_factor"] == pytest.approx(rows["L1"]["route_length_m"] * 1.1, abs=0.01)
    q = s.export_quantities("demo")
    assert q["named_range"] == "Quantities" and q["floors"] == ["4F"]
    wb = openpyxl.load_workbook(q["file"])
    assert "Quantities" in wb.defined_names
    assert any(r[1] == "luminaire" and r[3] >= 1 for r in wb["Quantities"].iter_rows(min_row=2, values_only=True))
    m = s.list_devices("demo", "4F")
    assert m["revision"] > 0
    with pytest.raises(ValueError, match="units are not known"):
        mm = dict(Store(server.CFG.workspace).load("demo", "4F"))
        mm["units"] = None
        tables.circuit_lengths(mm, 1.1, None)


# ------------------------------------------------------------------ live: tables, lengths, recovery
def test_live_table_import_batches_verify_and_lengths(s, live, tmp_path):
    sess, app = live
    table = filled_table(s, TABLE)
    dry = s.live_place_from_table("GF", table)
    assert dry["dry_run"] and dry["this_batch"] == 5 and not list(app.ActiveDocument.ModelSpace)
    s.live_place_from_table("GF", table, limit=3, dry_run=False)
    assert len(app.ActiveDocument._marks) == 1
    part = s.live_verify_placement("GF", table)
    assert not part["match"] and part["actual_total"] == 3
    done = s.live_place_from_table("GF", table, offset=3, limit=3, dry_run=False)
    assert done["remaining_after_this_batch"] == 0 and done["created"] == 2
    assert s.live_verify_placement("GF", table)["match"]
    s.live_assign("GF-SKT-01", "GF-P01", "DB-GF")
    s.live_place("GF", "db", 100, 4900)
    s.live_route("GF-P01", "power", db="DB-GF")
    lengths = s.live_circuit_lengths(export_project="proj")
    row = next(r for r in lengths["rows"] if r["circuit"] == "GF-P01")
    assert row["route_length_m"] > 0 and row["length_m_with_factor"] == pytest.approx(row["route_length_m"] * 1.1, abs=0.01)
    assert "GF-P02" in lengths["circuits_without_route"]
    assert "Circuit_Lengths_Live" in openpyxl.load_workbook(lengths["file"]).defined_names
    log = s.work_log()
    ops = [e["op"] for e in log["entries"]]
    assert "place_many" in ops and "route" in ops


def test_live_resync_after_a_dropped_connection(s, live):
    sess, app = live
    s.live_place("GF", "socket", 100, 100)
    s.live_place("GF", "socket", 900, 900)
    for g in app.ActiveDocument.Groups:  # the drawing lost a device (crash / manual delete)
        if g.Name.endswith("GF-SKT-02"):
            for it in g:
                it.Delete()
            g.Delete()
    rep = s.live_resync()
    assert rep["registered_but_not_in_drawing"] == ["GF-SKT-02"] and "prune=true" in rep["next"]
    assert s.live_resync(prune=True)["pruned"] is True
    assert s.live_resync()["next"] == "Registry and drawing agree."


def lisp_handler_for(app):
    doc = app.ActiveDocument

    def ents(layers=None, kind=None):
        out = []
        for e in doc.ModelSpace:
            if e.deleted:
                continue
            if layers is not None and e.Layer.lower() not in layers:
                continue
            if kind and {"AcDbLine": "LINE", "AcDbMesh": "MESH", "AcDbText": "TEXT"}.get(e.ObjectName) != kind:
                continue
            out.append(e)
        return out

    def handler(cmd):
        quoted = re.findall(r'"([^"]*)"', cmd)
        layer = re.search(r'\(cons 8 "([^"]+)"\)', cmd)
        if "_.ERASE" in cmd:
            lays = {x.lower() for x in layer.group(1).split(",")}
            victims = ents(lays)
            for e in victims:
                e.Delete()
            return str(len(victims))
        if "_.SCALE" in cmd:
            return str(len(ents()))
        if "(cons 0 ty)" in cmd:
            types = [q for q in quoted if q.isupper() and q != layer.group(1)]
            return "".join(f"{t}={len(ents({layer.group(1).lower()}, t))};" for t in types)
        if "(cons 8 l)" in cmd:
            names = [q for q in quoted if q not in ("=", ";")]
            return "".join(f"{n}={len(ents({n.lower()}))};" for n in names)
        if layer:
            return str(len(ents({x.lower() for x in layer.group(1).split(",")})))
        return str(len(ents()))
    return handler


def test_live_inventory_and_sanitize_on_a_heavy_export(s, live):
    sess, app = live
    doc = app.ActiveDocument
    for name in ("DLX_OBJ", "DLX_CALC", "Wall", "0"):
        if name not in [x.Name for x in doc.Layers]:
            doc.Layers.append(Layer(name))
    doc.Blocks.append(Block("LUM_A", 3))
    for _ in range(40):
        doc.ModelSpace.append(Ent("AcDbMesh", "DLX_OBJ"))
    for _ in range(5):
        doc.ModelSpace.append(Ent("AcDbLine", "Wall", StartPoint=(0, 0, 0), EndPoint=(5000, 0, 0)))
    doc.ModelSpace.append(Ent("AcDbLine", "DLX_CALC", StartPoint=(0, 0, 0), EndPoint=(1, 0, 0)))
    doc.lisp_handler = lisp_handler_for(app)

    inv = s.live_inventory()
    layers = {x["layer"]: x["entities"] for x in inv["layers"]}
    assert layers == {"DLX_OBJ": 40, "Wall": 5, "DLX_CALC": 1} and inv["entities"] == 46
    assert inv["top_layers_by_type"]["DLX_OBJ"] == {"MESH": 40}
    assert inv["blocks"] == [{"block": "LUM_A", "entities_in_definition": 3}]
    assert any("sanitised COPY" in n for n in inv["notes"])

    dry = s.live_sanitize(["DLX_OBJ", "DLX_CALC"], scale=1000)
    assert dry["dry_run"] and dry["to_erase"] == 41 and dry["entities_before"] == 46
    with pytest.raises(ValueError, match="confirm_copy"):
        s.live_sanitize(["DLX_OBJ"], dry_run=False)
    with pytest.raises(ValueError, match="Unsupported layer"):
        s.live_sanitize(['A"B'], dry_run=True)
    done = s.live_sanitize(["DLX_OBJ", "DLX_CALC"], scale=1000, dry_run=False, confirm_copy=True)
    assert done["entities_after"] == 5 and done["still_on_stripped_layers"] == 0
    assert any("_.SCALE" in c and "1000" in c for c in doc.commands)
    assert doc.vars["INSUNITS"] == 4
    s.base_register(doc.FullName)
    with pytest.raises(ValueError, match="registered BASE"):
        s.live_sanitize(["DLX_OBJ"])
