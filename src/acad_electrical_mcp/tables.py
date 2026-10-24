"""Table-driven drawing: placement from a sheet, verification against it, circuit lengths and
quantities exported to Excel as named ranges (so reports can link to them)."""

from __future__ import annotations

import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from . import workbook
from .dialux import UNIT_TO_M


def rows_for_floor(path: str, floor: str) -> list[dict[str, Any]]:
    res = workbook.read(path, "placement_table")
    if not res["ok"]:
        raise ValueError(f"The placement table has {res['errors']} error(s): "
                         f"{[(i['row'], i['column'], i['message']) for i in res['issues'][:5]]}")
    rows = [r for r in res["rows"] if str(r["floor"]).lower() == floor.lower()]
    if not rows:
        raise ValueError(f"No rows for floor {floor!r} in the placement table.")
    return rows


def resolve_room(model: dict[str, Any] | None, label: str) -> tuple[str, bool]:
    """Room id for a table's room cell (id or name). Returns (value, found_in_model)."""
    if not model:
        return label, True
    for r in model["rooms"]:
        if label.lower() in (r["id"].lower(), r["name"].lower()):
            return r["id"], True
    return label, False


def identical_rooms(rows: list[dict[str, Any]]) -> list[list[str]]:
    """Rooms whose device mix is identical (3+ devices) and not explained in Notes: usually a
    template copied from room to room."""
    by_room: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        by_room[str(r["room"])].append(r)
    sig: dict[tuple, list[str]] = defaultdict(list)
    for room, rs in by_room.items():
        if len(rs) >= 3 and not all(r.get("notes") for r in rs):
            sig[tuple(sorted(Counter(r["type"] for r in rs).items()))].append(room)
    return [v for v in sig.values() if len(v) > 1]


def specs_from_rows(model: dict[str, Any] | None, rows: list[dict[str, Any]]
                    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    specs, unknown_rooms, outside = [], set(), 0
    for r in rows:
        room, found = resolve_room(model, str(r["room"]))
        if not found:
            unknown_rooms.add(room)
        elif model:
            rm = next((x for x in model["rooms"] if x["id"] == room), None)
            if rm and not (rm["bounds"][0] <= r["x"] <= rm["bounds"][2]
                           and rm["bounds"][1] <= r["y"] <= rm["bounds"][3]):
                outside += 1
        spec = {"type": r["type"], "x": r["x"], "y": r["y"], "rotation": r.get("rotation") or 0.0,
                "room": room if found else None, "circuit": r.get("circuit"), "db": r.get("db"),
                "note": r.get("label") or r.get("notes")}
        specs.append(spec)
    return specs, {"unknown_rooms": sorted(unknown_rooms), "outside_their_room": outside}


def compare_counts(rows: list[dict[str, Any]], devices: list[dict[str, Any]],
                   model: dict[str, Any] | None = None) -> dict[str, Any]:
    """Expected (table) versus actual (drawing/model) device counts per room and type."""
    exp: Counter = Counter()
    for r in rows:
        room, _ = resolve_room(model, str(r["room"]))
        exp[(room, r["type"])] += 1
    act: Counter = Counter((d.get("room") or "(no room)", d["type"]) for d in devices
                           if d["type"] != "db" or any(t[1] == "db" for t in exp))
    keys = sorted(set(exp) | set(act))
    table = [{"room": k[0], "type": k[1], "expected": exp.get(k, 0), "actual": act.get(k, 0),
              "match": exp.get(k, 0) == act.get(k, 0)} for k in keys]
    bad = [t for t in table if not t["match"]]
    return {"match": not bad, "rows": table, "mismatches": bad,
            "expected_total": sum(exp.values()), "actual_total": sum(act.values())}


def route_length(points: list[list[float]]) -> float:
    return sum(math.hypot(b[0] - a[0], b[1] - a[1]) for a, b in zip(points, points[1:], strict=False))


def circuit_lengths(model: dict[str, Any], factor: float, unit_to_m: float | None) -> dict[str, Any]:
    u = unit_to_m or UNIT_TO_M.get(str(model.get("units")))
    if not u:
        raise ValueError("The drawing units are not known: pass unit_to_m (metres per drawing unit, "
                         "0.001 for a millimetre drawing).")
    if factor <= 0:
        raise ValueError("factor must be positive")
    routes = {r["circuit"]: r for r in model["routes"]}
    members: dict[str, int] = Counter(d["circuit"] for d in model["devices"] if d.get("circuit"))
    rows = []
    for c in sorted(model["circuits"], key=lambda c: c["id"]):
        r = routes.get(c["id"])
        ln = route_length(r["points"]) if r else None
        rows.append({"circuit": c["id"], "floor": model["floor"], "db": c.get("db"), "kind": c["kind"],
                     "devices": members.get(c["id"], 0),
                     "route_length_m": round(ln * u, 2) if ln else None,
                     "length_m_with_factor": round(ln * u * factor, 2) if ln else None,
                     "source": f"floor model revision {model['revision']}" if r else "NO ROUTE"})
    return {"unit_to_m": u, "factor": factor, "rows": rows,
            "circuits_without_route": [r["circuit"] for r in rows if r["source"] == "NO ROUTE"]}


def quantities(models: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for m in models:
        c = Counter(d["type"] for d in m["devices"])
        for t, n in sorted(c.items()):
            rows.append({"floor": m["floor"], "item": t, "detail": "", "qty": n, "unit": "no."})
        lum = Counter((d.get("luminaire_type") or "", d.get("watts")) for d in m["devices"]
                      if d["type"] == "luminaire" and d.get("luminaire_type"))
        for (lt, w), n in sorted(lum.items(), key=lambda kv: str(kv[0])):
            rows.append({"floor": m["floor"], "item": "luminaire type", "detail":
                         f"{lt}" + (f", {w} W" if w else ""), "qty": n, "unit": "no."})
    return rows


def write_xlsx(path: Path, sheet: str, headers: list[str], rows: list[list[Any]], range_name: str,
               title: str | None = None) -> str:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[get_column_letter(c)].width = max(14, len(h) + 4)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ref = f"'{sheet}'!$A$1:${get_column_letter(len(headers))}${max(len(rows) + 1, 2)}"
    wb.defined_names[range_name] = DefinedName(range_name, attr_text=ref)
    ws.freeze_panes = "A2"
    if title:
        wb.properties.title = title
    wb.save(str(path))
    return ref
